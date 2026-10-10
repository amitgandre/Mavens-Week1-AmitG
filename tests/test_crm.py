import io
from pathlib import Path
from unittest.mock import MagicMock, patch
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest
import crm

@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(crm, "DATA", tmp_path)
    monkeypatch.setattr(crm, "SMTP_CONFIG", tmp_path / ".env.smtp")


def account(username="seller", role="Salesperson", temporary=False):
    crm.create_account("Test Person", username, "test@example.com", role, "a-strong-password", temporary)


def test_login_role_hash_and_throttle():
    account()
    stored = crm.read_accounts()["seller"]["password_hash"]
    assert "a-strong-password" not in stored
    assert crm.authenticate("seller", "a-strong-password", "Salesperson")[0]
    assert not crm.authenticate("seller", "a-strong-password", "Executive Dashboard")[0]
    for _ in range(4):
        crm.authenticate("seller", "bad", "Salesperson")
    assert "15 minutes" in crm.authenticate("seller", "a-strong-password", "Salesperson")[1]


def test_excel_roundtrip_conflict_and_isolation():
    account()
    account("seller2")
    df = crm.sample_data()
    crm.save_workbook("seller", df, "")
    loaded, version = crm.load_workbook("seller")
    assert loaded.Amount.sum() == df.Amount.sum()
    loaded.loc[0, "Amount"] = 999
    crm.save_workbook("seller", loaded, version)
    assert crm.load_workbook("seller")[0].loc[0, "Amount"] == 999
    with pytest.raises(ValueError, match="another session"):
        crm.save_workbook("seller", df, version)
    assert crm.load_workbook("seller2")[0].empty
    assert len(crm.all_opportunities()) == 4


def test_invalid_input_and_formula_text():
    df = crm.sample_data()
    df.loc[0, "Probability"] = 101
    with pytest.raises(ValueError):
        crm.normalize(df)
    df = crm.sample_data()
    df.loc[0, "Notes"] = '=HYPERLINK("https://example.com")'
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(crm.excel_bytes(df)))
    assert wb["Opportunities"]["H2"].data_type == "s"
    with pytest.raises(ValueError):
        crm.workbook_path("../../etc")


def test_email_two_messages_and_password_change(monkeypatch):
    account(temporary=True)
    for key in ["SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "SMTP_FROM"]:
        monkeypatch.setenv(key, "test@example.com")
    with patch("crm.smtplib.SMTP") as mock:
        smtp = mock.return_value.__enter__.return_value
        crm.email_credentials("seller", "temporary-password")
        messages = [call.args[0] for call in smtp.send_message.call_args_list]
        assert len(messages) == 2
        assert "temporary-password" not in messages[0].get_content()
        assert "temporary-password" in messages[1].get_content()
        smtp.starttls.assert_called_once()
    crm.change_password("seller", "a-strong-password", "new-strong-password")
    assert not crm.read_accounts()["seller"]["temporary"]
    assert crm.authenticate("seller", "new-strong-password", "Salesperson")[0]


def test_app_roles_and_temporary_password():
    account()
    account("executive", "Executive Dashboard")
    account("admin", "Company Administrator")
    account("temporary", temporary=True)
    crm.save_workbook("seller", crm.sample_data(), "")
    for username in ["seller", "executive", "admin", "temporary"]:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()[username], username=username)
        app.run()
        assert not app.exception, [e.message for e in app.exception]
        if username == "temporary":
            app.text_input[0].set_value("a-strong-password")
            app.text_input[1].set_value("new-strong-password")
            app.text_input[2].set_value("new-strong-password")
            next(b for b in app.button if b.label == "Update password").click()
            app.run()
            assert not app.exception
            assert not app.session_state["user"]["temporary"]


def test_smtp_file_loaded_and_reloaded(monkeypatch):
    for key in crm.SMTP_KEYS:
        monkeypatch.delenv(key, raising=False)
    assert not crm.smtp_ready()
    crm.SMTP_CONFIG.write_text('SMTP_HOST="smtp.gmail.com"\nSMTP_USER="test@example.com"\nSMTP_PASSWORD="fake-test-secret"\nSMTP_FROM="test@example.com"\n')
    assert crm.smtp_ready()
    assert crm.smtp_settings()["SMTP_HOST"] == "smtp.gmail.com"
    crm.SMTP_CONFIG.write_text('SMTP_HOST="smtp.gmail.com"\n')
    assert not crm.smtp_ready()
    assert "SMTP_PASSWORD" in crm.smtp_missing()


def test_admin_password_validation_permissions_and_login():
    account()
    account("admin", "Company Administrator")
    for _ in range(5):
        crm.authenticate("seller", "wrong", "Salesperson")
    with pytest.raises(ValueError, match="12 characters"):
        crm.set_user_password("admin", "seller", "short")
    with pytest.raises(ValueError, match="administrator"):
        crm.set_user_password("seller", "admin", "new-chosen-password")
    assert crm.verify_password("a-strong-password", crm.read_accounts()["seller"]["password_hash"])
    crm.set_user_password("admin", "seller", "admin-chosen-password")
    user, error = crm.authenticate("seller", "admin-chosen-password", "Salesperson")
    assert user and user["temporary"] and error is None
    assert not crm.verify_password("a-strong-password", user["password_hash"])
    assert "admin-chosen-password" not in (crm.DATA / "users.json").read_text()
    crm.update_account("admin", active=False)
    with pytest.raises(ValueError, match="administrator"):
        crm.set_user_password("admin", "seller", "another-chosen-password")


def test_admin_chosen_password_ui(monkeypatch):
    account("admin", "Company Administrator")
    account()
    monkeypatch.setattr(crm, "smtp_ready", lambda: True)
    with patch("crm.email_credentials") as send:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()["admin"], username="admin")
        app.run()
        def fill(label, value):
            next(field for field in app.text_input if field.label == label).set_value(value)
        fill("Full name", "New Seller")
        fill("Email", "new@example.com")
        fill("Username", "new-seller")
        fill("Temporary password (12+ characters)", "chosen-create-password")
        fill("Confirm temporary password", "chosen-create-password")
        next(b for b in app.button if b.label == "Create account").click()
        app.run()
        assert not app.exception
        send.assert_called_once_with("new-seller", "chosen-create-password")
        assert crm.verify_password("chosen-create-password", crm.read_accounts()["new-seller"]["password_hash"])
        next(field for field in app.selectbox if field.label == "User").select("seller")
        fill("New temporary password (12+ characters)", "chosen-reset-password")
        fill("Confirm new temporary password", "chosen-reset-password")
        next(b for b in app.button if b.label == "Set user password").click()
        app.run()
        assert not app.exception
        send.assert_called_with("seller", "chosen-reset-password")
        assert crm.verify_password("chosen-reset-password", crm.read_accounts()["seller"]["password_hash"])


def pipeline_data():
    return pd.DataFrame([
        ["001", "Acme", "Expansion", 100000, "2026-01-01", "2026-12-01", 75, "Open", 300000, 100000, "3 years", 25, "Technology", ""],
        ["002", "Beta", "Renewal", 25000, "2026-01-01", "2026-12-01", 50, "Closed", None, None, "12 months", None, "Services", "seller"],
    ], columns=crm.PIPELINE_COLUMNS)


def test_pipeline_template_import_edit_export_and_ownership():
    from openpyxl import load_workbook
    wb = load_workbook(crm.PIPELINE_TEMPLATE)
    assert list(next(wb["Sheet1"].values)) == crm.PIPELINE_COLUMNS
    df = crm.read_upload(crm.excel_bytes(pipeline_data()), "seller")
    assert list(df.columns) == crm.PIPELINE_COLUMNS
    assert df["Opportunity ID"].tolist() == ["001", "002"]
    assert df["Salesperson Username"].tolist() == ["seller", "seller"]
    crm.save_workbook("seller", df, "")
    saved, version = crm.load_workbook("seller")
    saved.loc[0, "ACV"] = 120000
    saved.loc[0, "Contract Term"] = "4 years"
    crm.save_workbook("seller", saved, version)
    reloaded, _ = crm.load_workbook("seller")
    assert reloaded.loc[0, "ACV"] == 120000
    assert reloaded.loc[0, "TCV"] == 300000
    assert reloaded.loc[0, "Contract Term"] == "4 years"
    assert reloaded.loc[0, "Margin"] == 25
    assert reloaded.loc[0, "Opportunity ID"] == "001"
    charts = crm.analytics_data(reloaded)
    assert charts.loc[0, "Amount"] == 100000
    assert charts.loc[1, "Stage"] == "Closed"
    with pytest.raises(ValueError, match="signed-in username"):
        crm.read_upload(crm.excel_bytes(pipeline_data()), "other-seller")
    empty = crm.read_upload(crm.PIPELINE_TEMPLATE.read_bytes(), "seller")
    assert empty.empty and list(empty.columns) == crm.PIPELINE_COLUMNS


def test_pipeline_validation_and_dashboard():
    invalid = pipeline_data()
    invalid.loc[1, "Opportunity ID"] = "001"
    with pytest.raises(ValueError, match="unique"):
        crm.normalize(invalid)
    account()
    account("executive", "Executive Dashboard")
    crm.save_workbook("seller", pipeline_data(), "")
    for username in ["seller", "executive"]:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()[username], username=username)
        app.run()
        assert not app.exception, [e.message for e in app.exception]
        assert next(m for m in app.metric if m.label == "Open pipeline").value == "$0.10M"
    crm.save_workbook("empty-seller", crm.read_upload(crm.PIPELINE_TEMPLATE.read_bytes(), "empty-seller"), "")
    assert crm.load_workbook("empty-seller")[0].empty


def test_salesperson_table_columns_and_legacy_backup():
    account()
    legacy = crm.sample_data()
    crm.save_workbook("seller", legacy, "")
    table, version = crm.load_salesperson_table("seller")
    assert list(table.columns) == crm.PIPELINE_COLUMNS
    assert table["Opportunity Value"].sum() == legacy.Amount.sum()
    assert table["Salesperson Username"].eq("seller").all()
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["seller"], username="seller")
    app.session_state["sheet"] = legacy
    app.session_state["version"] = version
    app.run()
    assert not app.exception
    assert list(app.dataframe[0].value.columns) == ["Select"] + crm.PIPELINE_COLUMNS
    crm.save_workbook("seller", table, version)
    backup = crm.DATA / "opportunities_seller.legacy-backup.xlsx"
    assert backup.exists()
    assert pd.read_excel(backup)["Notes"].tolist() == legacy.Notes.tolist()
    assert list(crm.load_workbook("seller")[0].columns) == crm.PIPELINE_COLUMNS


def test_salesperson_add_delete_rows_draft_and_save():
    account()
    crm.save_workbook("seller", pipeline_data(), "")
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["seller"], username="seller")
    app.run()
    next(b for b in app.button if b.label == "Add new row").click()
    app.run()
    assert not app.exception
    assert len(app.dataframe[0].value) == 3
    assert len(crm.load_workbook("seller")[0]) == 2
    added = app.dataframe[0].value.iloc[2]
    assert added["Opportunity ID"].startswith("OPP-")
    assert added["Salesperson Username"] == "seller"
    generation = app.session_state["editor_generation"]
    app.session_state[f"million_pipeline_editor_{generation}"] = {"edited_rows": {0: {"Select": True}, 2: {"Select": True}}, "added_rows": [], "deleted_rows": []}
    next(b for b in app.button if b.label == "Delete selected rows").click()
    app.run()
    assert not app.exception
    assert app.dataframe[0].value["Opportunity ID"].tolist() == ["002"]
    assert len(crm.load_workbook("seller")[0]) == 2
    next(b for b in app.button if b.label == "Save changes to Excel").click()
    app.run()
    assert not app.exception
    assert crm.load_workbook("seller")[0]["Opportunity ID"].tolist() == ["002"]
    next(b for b in app.button if b.label == "Select all rows").click()
    app.run()
    assert app.dataframe[0].value["Select"].all()
    next(b for b in app.button if b.label == "Delete selected rows").click()
    app.run()
    next(b for b in app.button if b.label == "Save changes to Excel").click()
    app.run()
    assert not app.exception
    assert crm.load_workbook("seller")[0].empty


def test_select_all_clear_selection_preserves_edits():
    account()
    crm.save_workbook("seller", pipeline_data(), "")
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["seller"], username="seller")
    app.run()
    app.session_state["million_pipeline_editor_0"] = {"edited_rows": {0: {"Opportunity Value": 0.123456}}, "added_rows": [], "deleted_rows": []}
    next(b for b in app.button if b.label == "Select all rows").click()
    app.run()
    assert not app.exception
    assert app.dataframe[0].value["Select"].all()
    assert app.dataframe[0].value.iloc[0]["Opportunity Value"] == 0.123456
    next(b for b in app.button if b.label == "Clear selection").click()
    app.run()
    assert not app.dataframe[0].value["Select"].any()
    assert app.dataframe[0].value.iloc[0]["Opportunity Value"] == 0.123456
    next(b for b in app.button if b.label == "Save changes to Excel").click()
    app.run()
    saved, _ = crm.load_workbook("seller")
    assert "Select" not in saved.columns
    assert saved.iloc[0]["Opportunity Value"] == 123456


def test_subset_updates_preserve_rows_ownership_and_conflicts():
    account()
    account("seller2")
    account("executive", "Executive Dashboard")
    df = pipeline_data()
    crm.save_workbook("seller", df, "")
    second = df.copy()
    second["Salesperson Username"] = "seller2"
    crm.save_workbook("seller2", second, "")
    all_rows = crm.all_opportunities()
    subset = all_rows[all_rows["Opportunity ID"] == "001"].copy().iloc[::-1]
    edited = crm.visible_opportunities(subset)
    assert not set(["Amount", "Customer", "Stage", "Probability", "Salesperson", "__row", "__version"]) & set(edited.columns)
    edited.loc[0, "Opportunity Value"] = 88888
    edited.loc[1, "Margin"] = 35
    with pytest.raises(ValueError, match="own opportunities"):
        crm.update_opportunity_rows("seller", subset, edited)
    assert crm.update_opportunity_rows("executive", subset, edited) == 2
    assert crm.load_workbook("seller2")[0].loc[0, "Opportunity Value"] == 88888
    assert crm.load_workbook("seller")[0].loc[0, "Margin"] == 35
    assert crm.load_workbook("seller2")[0].loc[1, "Opportunity Value"] == 25000
    with pytest.raises(ValueError, match="changed in another session"):
        crm.update_opportunity_rows("executive", subset, edited)


def test_subset_validation_before_any_workbook_write():
    account()
    account("seller2")
    account("executive", "Executive Dashboard")
    crm.save_workbook("seller", pipeline_data(), "")
    second = pipeline_data()
    second["Salesperson Username"] = "seller2"
    crm.save_workbook("seller2", second, "")
    original = crm.all_opportunities()
    edited = crm.visible_opportunities(original)
    edited.loc[0, "Opportunity Value"] = 123
    edited.loc[2, "Win Probability"] = 101
    with pytest.raises(ValueError):
        crm.update_opportunity_rows("executive", original, edited)
    assert crm.load_workbook("seller")[0].loc[0, "Opportunity Value"] == 100000


def test_edit_intelligence_and_executive_tables():
    account()
    account("executive", "Executive Dashboard")
    crm.save_workbook("seller", pipeline_data(), "")
    for username, table_name, label, value in [
        ("seller", "intelligence", "Save intelligence changes to Excel", 77777),
        ("executive", "all", "Save all opportunity changes to Excel", 99999),
    ]:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()[username], username=username)
        app.run()
        assert not app.exception
        for table in app.dataframe:
            columns = list(table.value.columns)
            assert len(columns) == len(set(columns))
            assert "Amount" not in columns and "Probability" not in columns
        snapshot = next(key for key in app.session_state._state.filtered_state if key.startswith(f"snapshot_{table_name}_"))
        generation = snapshot.removeprefix("snapshot_")
        app.session_state[f"million_table_{generation}"] = {"edited_rows": {0: {"Opportunity Value": value / 1_000_000}}, "added_rows": [], "deleted_rows": []}
        next(b for b in app.button if b.label == label).click()
        app.run()
        assert not app.exception
        saved, _ = crm.load_workbook("seller")
        assert saved.loc[0, "Opportunity Value"] == value
        assert saved.loc[1, "Opportunity Value"] == 25000


def test_executive_empty_company_table():
    account("executive", "Executive Dashboard")
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["executive"], username="executive")
    app.run()
    assert not app.exception
    assert list(app.dataframe[0].value.columns) == crm.PIPELINE_COLUMNS


def test_account_role_change_delete_and_authorization():
    account("admin", "Company Administrator")
    account()
    crm.save_workbook("seller", pipeline_data(), "")
    with pytest.raises(ValueError, match="administrator"):
        crm.manage_account("seller", "admin", delete=True)
    with pytest.raises(ValueError, match="own"):
        crm.manage_account("admin", "admin", delete=True)
    with pytest.raises(ValueError, match="supported role"):
        crm.manage_account("admin", "seller", role="Invalid")
    crm.manage_account("admin", "seller", role="Executive Dashboard")
    assert not crm.authenticate("seller", "a-strong-password", "Salesperson")[0]
    assert crm.authenticate("seller", "a-strong-password", "Executive Dashboard")[0]
    crm.manage_account("admin", "seller", delete=True)
    assert "seller" not in crm.read_accounts()
    assert not crm.authenticate("seller", "a-strong-password", "Executive Dashboard")[0]
    assert not crm.workbook_path("seller").exists()
    assert crm.archive_path("seller").exists()
    assert crm.all_opportunities().empty
    with pytest.raises(ValueError, match="already exists"):
        account()
    account("another")
    assert crm.read_accounts(include_deleted=True)["seller"]["deleted"]


def test_admin_account_controls_ui():
    account("admin", "Company Administrator")
    account()
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["admin"], username="admin")
    app.run()
    next(field for field in app.selectbox if field.label == "New role").select("Executive Dashboard")
    next(button for button in app.button if button.label == "Change role").click()
    app.run()
    assert not app.exception
    assert crm.read_accounts()["seller"]["role"] == "Executive Dashboard"
    assert next(button for button in app.button if button.label == "Delete account").disabled
    next(field for field in app.checkbox if field.label == "Confirm deletion of seller").check()
    app.run()
    next(button for button in app.button if button.label == "Delete account").click()
    app.run()
    assert not app.exception
    assert "seller" not in crm.read_accounts()


def test_pipeline_archive_exact_bytes_permissions_and_ui():
    account("admin", "Company Administrator")
    account()
    crm.save_workbook("seller", pipeline_data(), "")
    original_bytes = crm.workbook_path("seller").read_bytes()
    crm.manage_account("admin", "seller", delete=True)
    assert not crm.workbook_path("seller").exists()
    records = crm.archived_pipelines("admin")
    assert [row["Username"] for row in records] == ["seller"]
    table, content = crm.read_archived_pipeline("admin", "seller")
    assert content == original_bytes
    assert len(table) == 2
    assert list(table.columns) == crm.PIPELINE_COLUMNS
    with pytest.raises(ValueError, match="administrator"):
        crm.read_archived_pipeline("seller", "seller")
    with pytest.raises(ValueError, match="deleted"):
        crm.save_workbook("seller", pipeline_data(), "")
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["admin"], username="admin")
    app.run()
    assert not app.exception
    assert next(field for field in app.selectbox if field.label == "Archived salesperson").value == "seller"
    assert any(len(frame.value) == 2 and "Opportunity Value" in frame.value for frame in app.dataframe)


def test_archive_previously_deleted_accounts_and_missing_pipeline():
    account("admin", "Company Administrator")
    account()
    crm.save_workbook("seller", pipeline_data(), "")
    crm.update_account("seller", active=False, deleted=True)
    assert crm.archived_pipelines("admin")[0]["Username"] == "seller"
    assert not crm.workbook_path("seller").exists()
    assert len(crm.archived_pipelines("admin")) == 1
    account("empty-seller")
    crm.manage_account("admin", "empty-seller", delete=True)
    assert len(crm.archived_pipelines("admin")) == 1


def test_smtp_authentication_errors_are_safe_and_recorded(monkeypatch):
    import smtplib
    import ssl
    account()
    for key in ["SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "SMTP_FROM"]:
        monkeypatch.setenv(key, "fake-test-value")
    with patch("crm.smtplib.SMTP") as mock:
        smtp = mock.return_value.__enter__.return_value
        smtp.login.side_effect = smtplib.SMTPAuthenticationError(535, b"private-server-text")
        with pytest.raises(ValueError, match="SMTP 535") as error:
            crm.email_credentials("seller", "temporary-password")
        assert "private-server-text" not in str(error.value)
        assert "temporary-password" not in str(error.value)
        assert "0/2" in crm.read_accounts()["seller"]["email_status"]
        smtp.send_message.assert_not_called()
        context = smtp.starttls.call_args.kwargs["context"]
        assert context.verify_mode == ssl.CERT_REQUIRED and context.check_hostname
    assert "closed the connection" in crm.smtp_error_message(smtplib.SMTPServerDisconnected("sensitive"))


def test_gmail_password_spaces_and_check_without_email(monkeypatch):
    monkeypatch.setenv("SMTP_HOST", "smtp.gmail.com")
    monkeypatch.setenv("SMTP_USER", "test@example.com")
    monkeypatch.setenv("SMTP_FROM", "test@example.com")
    monkeypatch.setenv("SMTP_PASSWORD", "abcd efgh ijkl mnop")
    assert crm.smtp_settings()["SMTP_PASSWORD"] == "abcdefghijklmnop"
    with patch("crm.smtplib.SMTP") as mock:
        crm.check_smtp()
        smtp = mock.return_value.__enter__.return_value
        smtp.login.assert_called_once_with("test@example.com", "abcdefghijklmnop")
        smtp.send_message.assert_not_called()


def test_smtp_port_465_uses_verified_tls_without_starttls(monkeypatch):
    monkeypatch.setenv("SMTP_HOST", "smtp.gmail.com")
    monkeypatch.setenv("SMTP_PORT", "465")
    for key in ["SMTP_USER", "SMTP_FROM", "SMTP_PASSWORD"]:
        monkeypatch.setenv(key, "fake-test-value")
    with patch("crm.smtplib.SMTP_SSL") as mock:
        crm.check_smtp()
        smtp = mock.return_value.__enter__.return_value
        smtp.starttls.assert_not_called()
        smtp.login.assert_called_once()
        assert mock.call_args.kwargs["context"].check_hostname
        smtp.send_message.assert_not_called()


def test_setup_reports_authentication_disconnect_without_saving(monkeypatch):
    import configure_smtp
    import smtplib
    monkeypatch.setattr("sys.argv", ["configure_smtp.py"])
    monkeypatch.setattr("builtins.input", lambda _: "test@example.com")
    monkeypatch.setattr("getpass.getpass", lambda _: "abcdefghijklmnop")
    with patch("configure_smtp.smtplib.SMTP") as mock:
        mock.return_value.__enter__.return_value.login.side_effect = smtplib.SMTPServerDisconnected("private-text")
        with patch("configure_smtp.tempfile.mkstemp") as write:
            with pytest.raises(SystemExit, match="Gmail authentication failed") as error:
                configure_smtp.main()
            assert "private-text" not in str(error.value)
            assert "abcdefghijklmnop" not in str(error.value)
            write.assert_not_called()


def test_won_revenue_probability_and_closure_date():
    from datetime import datetime
    from zoneinfo import ZoneInfo
    today = pd.Timestamp(datetime.now(ZoneInfo("America/Chicago")).date())
    account()
    account("executive", "Executive Dashboard")
    rows = pd.concat([pipeline_data().iloc[:1]] * 4, ignore_index=True)
    rows["Opportunity ID"] = ["past", "today", "future", "partial"]
    rows["Opportunity Value"] = [1000, 2000, 4000, 8000]
    rows["Win Probability"] = [100, 100, 100, 99]
    rows["Open or Closed"] = ["Open", "Closed", "Open", "Open"]
    rows["Expected Closure Date"] = [today - pd.Timedelta(days=1), today, today + pd.Timedelta(days=1), today]
    crm.save_workbook("seller", rows, "")
    for username in ["seller", "executive"]:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()[username], username=username)
        app.run()
        assert not app.exception
        assert next(metric for metric in app.metric if metric.label == "Won revenue").value == "$0.00M"


def test_won_revenue_requires_nonzero_tcv_or_acv():
    from datetime import datetime
    from zoneinfo import ZoneInfo
    today = pd.Timestamp(datetime.now(ZoneInfo("America/Chicago")).date())
    account()
    account("executive", "Executive Dashboard")
    rows = pd.concat([pipeline_data().iloc[:1]] * 5, ignore_index=True)
    rows["Opportunity ID"] = ["tcv", "acv", "both-zero", "both-blank", "one-blank"]
    rows["Opportunity Value"] = [1000, 2000, 4000, 8000, 3000]
    rows["Win Probability"] = 100
    rows["Expected Closure Date"] = today
    rows["TCV"] = [100, 0, 0, None, None]
    rows["ACV"] = [0, 100, 0, None, 100]
    crm.save_workbook("seller", rows, "")
    for username in ["seller", "executive"]:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()[username], username=username)
        app.run()
        assert not app.exception
        assert next(metric for metric in app.metric if metric.label == "Won revenue").value == "$0.01M"


def test_account_email_reminders_and_reset(monkeypatch):
    import smtplib
    from contextlib import contextmanager
    account("admin", "Company Administrator")
    account()
    monkeypatch.setattr(crm, "smtp_ready", lambda: True)
    smtp = MagicMock()
    @contextmanager
    def connection(settings):
        yield smtp
    monkeypatch.setattr(crm, "smtp_connection", connection)
    previous_hash = crm.read_accounts()["seller"]["password_hash"]
    crm.send_account_reminder("admin", "seller")
    username_message = smtp.send_message.call_args.args[0]
    assert username_message["To"] == "test@example.com"
    assert "Username: seller" in username_message.get_content()
    assert "password" not in username_message.get_content().lower()
    assert crm.read_accounts()["seller"]["password_hash"] == previous_hash
    crm.send_account_reminder("admin", "seller", reset_password=True)
    password_message = smtp.send_message.call_args.args[0]
    password = password_message.get_content().split("Your temporary password: ")[1].splitlines()[0]
    assert len(password) == 12 and password.isascii() and password.isalnum()
    assert password_message["To"] == "test@example.com"
    assert "Username:" not in password_message.get_content()
    assert crm.verify_password(password, crm.read_accounts()["seller"]["password_hash"])
    assert crm.read_accounts()["seller"]["temporary"]
    assert password not in (crm.DATA / "users.json").read_text()
    with pytest.raises(ValueError, match="administrator"):
        crm.send_account_reminder("seller", "admin")
    current_hash = crm.read_accounts()["seller"]["password_hash"]
    @contextmanager
    def failure(settings):
        raise smtplib.SMTPAuthenticationError(535, b"private")
        yield
    monkeypatch.setattr(crm, "smtp_connection", failure)
    with pytest.raises(ValueError, match="authentication failed"):
        crm.send_account_reminder("admin", "seller", reset_password=True)
    assert crm.read_accounts()["seller"]["password_hash"] == current_hash


def test_account_row_email_buttons_ui(monkeypatch):
    account("admin", "Company Administrator")
    account()
    monkeypatch.setattr(crm, "smtp_ready", lambda: True)
    with patch("crm.send_account_reminder") as send:
        app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
        app.session_state["user"] = dict(crm.read_accounts()["admin"], username="admin")
        app.run()
        next(button for button in app.button if button.key == "resend_username_seller").click()
        app.run()
        assert not app.exception
        send.assert_called_with("admin", "seller", reset_password=False)
        next(button for button in app.button if button.key == "reset_email_seller").click()
        app.run()
        assert not app.exception
        send.assert_called_with("admin", "seller", reset_password=True)


def test_change_registered_email_and_reminder_recipient(monkeypatch):
    from contextlib import contextmanager
    account("admin", "Company Administrator")
    account()
    original = crm.read_accounts()["seller"]["password_hash"]
    with pytest.raises(ValueError, match="valid email"):
        crm.change_user_email("admin", "seller", "not-an-email")
    with pytest.raises(ValueError, match="administrator"):
        crm.change_user_email("seller", "admin", "new@example.com")
    crm.change_user_email("admin", "seller", " updated@example.com ")
    assert crm.read_accounts()["seller"]["email"] == "updated@example.com"
    assert crm.read_accounts()["seller"]["password_hash"] == original
    monkeypatch.setattr(crm, "smtp_ready", lambda: True)
    smtp = MagicMock()
    @contextmanager
    def connection(settings):
        yield smtp
    monkeypatch.setattr(crm, "smtp_connection", connection)
    crm.send_account_reminder("admin", "seller")
    assert smtp.send_message.call_args.args[0]["To"] == "updated@example.com"
    crm.manage_account("admin", "seller", delete=True)
    with pytest.raises(ValueError, match="does not exist"):
        crm.change_user_email("admin", "seller", "other@example.com")


def test_change_user_email_ui():
    account("admin", "Company Administrator")
    account()
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=20)
    app.session_state["user"] = dict(crm.read_accounts()["admin"], username="admin")
    app.run()
    next(field for field in app.selectbox if field.label == "Account for email change").select("seller")
    app.run()
    next(field for field in app.text_input if field.label == "New email address").set_value("updated@example.com")
    next(button for button in app.button if button.label == "Update email address").click()
    app.run()
    assert not app.exception
    assert crm.read_accounts()["seller"]["email"] == "updated@example.com"


def test_client_name_links_and_plain_names_saved():
    from urllib.parse import parse_qs, urlparse
    plain = pipeline_data()
    plain.loc[0, "Client Name"] = "Acme & Sons #1"
    linked = crm.linked_client_table(plain)
    url = linked.loc[0, "Client Name"]
    assert urlparse(url).netloc == "www.google.com"
    assert '"Acme & Sons #1"' in parse_qs(urlparse(url).query)["q"][0]
    assert "#client=Acme & Sons #1" in url
    restored = crm.restore_client_names(linked, plain)
    assert restored["Client Name"].tolist() == plain["Client Name"].tolist()
    linked.loc[0, "Client Name"] = "New Client"
    restored = crm.restore_client_names(linked, plain)
    assert restored.loc[0, "Client Name"] == "New Client"
    assert crm.client_research_url("") == ""
    account()
    crm.save_workbook("seller", restored, "")
    assert crm.load_workbook("seller")[0].loc[0, "Client Name"] == "New Client"
    assert "google.com" not in crm.load_workbook("seller")[0].loc[1, "Client Name"]


def test_persistent_session_restore_expiry_and_revocation():
    import json
    account()
    token = crm.create_session('seller')
    assert token not in (crm.DATA / 'sessions.json').read_text()
    assert crm.resume_session(token)['username'] == 'seller'
    assert crm.resume_session('invalid') is None
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'app.py', default_timeout=20)
    from unittest.mock import PropertyMock
    # A fresh browser connection restores authentication from its cookie.
    with patch('streamlit.runtime.context.ContextProxy.cookies', new_callable=PropertyMock, return_value={'crm_session': token}):
        app.run()
        assert not app.exception
        assert app.session_state['user']['username'] == 'seller'
        next(button for button in app.button if button.label == 'Sign out').click()
        app.run()
        assert not app.exception
        assert 'user' not in app.session_state
        assert crm.resume_session(token) is None
    crm.revoke_session(token)
    assert crm.resume_session(token) is None
    token = crm.create_session('seller')
    sessions = json.loads((crm.DATA / 'sessions.json').read_text())
    for record in sessions.values():
        record['expires'] = 0
    crm.atomic_write(crm.DATA / 'sessions.json', json.dumps(sessions).encode())
    assert crm.resume_session(token) is None


def test_persistent_session_invalidated_by_account_changes():
    account()
    token = crm.create_session('seller')
    crm.change_password('seller', 'a-strong-password', 'another-strong-password')
    assert crm.resume_session(token) is None
    token = crm.create_session('seller')
    crm.update_account('seller', role='Executive Dashboard')
    assert crm.resume_session(token) is None
    token = crm.create_session('seller')
    crm.update_account('seller', active=False)
    assert crm.resume_session(token) is None


def test_session_idle_timeout_and_activity(monkeypatch):
    account()
    now = [100000.0]
    monkeypatch.setattr(crm.time, 'time', lambda: now[0])
    token = crm.create_session('seller')
    now[0] += 899
    assert crm.resume_session(token)
    # Reading or refreshing a session does not extend its idle deadline.
    now[0] += 1
    assert crm.resume_session(token) is None
    assert crm.resume_session(token, touch=True) is None
    token = crm.create_session('seller')
    now[0] += 600
    assert crm.resume_session(token, touch=True, activity_at=now[0] - 20)
    now[0] += 879
    assert crm.resume_session(token)
    now[0] += 1
    assert crm.resume_session(token) is None


def test_browser_idle_event_signs_out():
    account()
    token = crm.create_session('seller')
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / 'app.py', default_timeout=20)
    app.session_state['user'] = crm.resume_session(token)
    app.session_state['session_token'] = token
    app.session_state['browser_session'] = {'token': token, 'kind': 'expired', 'id': 'timeout-test'}
    app.run()
    assert not app.exception
    assert 'user' not in app.session_state
    assert crm.resume_session(token) is None
    assert any('session expired' in message.value for message in app.info)
